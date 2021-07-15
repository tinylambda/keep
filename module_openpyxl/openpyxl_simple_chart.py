import pathlib

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference


if __name__ == '__main__':
    workbook = Workbook()
    sheet = workbook.create_sheet('Sample Sheet')
    rows = [
        ["Product", "Online", "Store"],
        [1, 30, 45],
        [2, 40, 30],
        [3, 40, 25],
        [4, 50, 30],
        [5, 30, 25],
        [6, 25, 35],
        [7, 20, 40],
    ]

    for row in rows:
        sheet.append(row)

    chart = BarChart()
    data = Reference(
        worksheet=sheet,
        min_row=1,
        max_row=8,
        min_col=2,
        max_col=3,
    )
    chart.add_data(data, titles_from_data=True)
    sheet.add_chart(chart, 'E2')

    filepath = pathlib.Path.home().joinpath('Downloads/test9.xlsx')
    workbook.save(filepath)

